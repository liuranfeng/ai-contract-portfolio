#!/usr/bin/env python3
"""Export an independent complete comparison report as DOCX or PDF.

The comparison JSON is authoritative. This exporter never infers semantic
equivalence, suppresses pending items, or modifies the original Word document.
"""
import argparse
import hashlib
import io
import json
import os
import re
import tempfile
from pathlib import Path
from xml.sax.saxutils import escape

OPERATIONS = {"equal": "一致", "insert": "新增", "delete": "删除", "replace": "修改"}
SEMANTICS = {"equal": "文字一致", "equivalent": "等义改写", "changed": "含义变化",
             "pending": "待判断", "uncertain": "需人工确认", "not_requested": "未请求语义判断"}


def _semantic_status(row, result):
    return row.get("semantic", {}).get("status", "pending" if result.get("mode") == "semantic" else "not_requested")


def _blocks(result):
    if str(result.get("schemaVersion")) != "1.0" or len(result.get("documents", [])) != 2:
        raise ValueError("Expected comparison schemaVersion 1.0 with two documents")
    alignments = result.get("alignments", [])
    pending = sum(_semantic_status(row, result) in ("pending", "uncertain") for row in alignments)
    out = [("title", "文档比对独立报告")]
    left, right = result["documents"]
    out += [("body", f"本报告记录 {left.get('name', '原件')} 与 {right.get('name', '比对件')} 的全部对齐条目、文字差异和语义判定。比较方向为原件到比对件。"),
            ("body", f"共 {len(alignments)} 个对齐条目，其中 {pending} 项待判断或需人工确认。报告是独立说明文件，不是可接受或拒绝修订的原始 Word。"),
            ("h1", "比对概况"),
            ("body", f"比较编号 {result.get('comparisonId', '')}"),
            ("body", f"校验方式 {'语义一致性校验' if result.get('mode') == 'semantic' else '字符一致性校验'}   文本范围 {'可阅读正文' if result.get('surface') == 'reading' else '解析原文'}   处理状态 {'需要复核' if result.get('status') == 'needs_review' else '已完成'}"),
            ("summary", [(OPERATIONS[op], sum(row.get('operation') == op for row in alignments)) for op in ("equal", "insert", "delete", "replace")] + [("待判断或人工确认", pending)])]
    for warning in result.get("warnings", []):
        out.append(("body", "注意 " + str(warning)))
    out.append(("h1", "全部比对条目"))
    for item in alignments:
        out.append(("h2", f"{item.get('id', '')} {OPERATIONS.get(item.get('operation'), '未分类')}"))
        if item.get("alignmentStatus") == "unresolved":
            out.append(("body", "对齐状态 待确认匹配关系"))
        for side, label in (("left", "原件"), ("right", "比对件")):
            refs = item.get(side + "Refs", [])
            identifiers = "；".join(f"documentIndex={r.get('documentIndex')} layoutId={json.dumps(r.get('layoutId'), ensure_ascii=False)}" for r in refs)
            out.append(("body", f"{label}来源 {identifiers or '空侧 无来源条目'}"))
            value = item.get(side + "Text", "")
            out.append(("label", label + "全文"))
            out.append(("quote", value if value else "空文本"))
        semantic = item.get("semantic", {})
        status = _semantic_status(item, result)
        out.append(("body", "语义状态 " + SEMANTICS.get(status, str(status))))
        if status == "equal":
            out.append(("body", "两侧文字相同，无需额外语义判断。"))
        elif status == "not_requested":
            out.append(("body", "本次只校验字符差异。"))
        else:
            reason = semantic.get("reason") or "待结合原文及上下文提供判定理由。"
            out.append(("body", "语义理由 " + reason))
        for key, label in (("leftQuote", "原件证据引文"), ("rightQuote", "比对件证据引文")):
            if semantic.get(key):
                out.append(("body", label + " " + semantic[key]))
    out.append(("h1", "来源文档索引"))
    for ordinal, doc in enumerate(result["documents"]):
        out.append(("h2", f"文档 {doc.get('documentIndex', ordinal)} {doc.get('name', '')}"))
        for key in ("documentIndex", "parseRunId", "sourceSha256", "rawParseSha256"):
            out.append(("body", f"{key} {doc.get(key) if doc.get(key) is not None else '未提供'}"))
        for layout in doc.get("layouts", []):
            out.append(("body", f"layoutId={json.dumps(layout.get('layoutId'), ensure_ascii=False)}  ordinal={layout.get('ordinal', '')}  page={layout.get('page', '')}  type={layout.get('type', '')}"))
    changed_ranges = [(item, [op for op in item.get("charOps", []) if op.get("op") != "equal"]) for item in alignments]
    if any(ops for _, ops in changed_ranges):
        out.append(("h1", "字符范围追溯"))
        out.append(("body", "下列范围采用 Unicode 码点零基半开区间，用于回查机器结果。"))
        for item, ops in changed_ranges:
            for op in ops:
                out.append(("body", f"{item.get('id')} {OPERATIONS.get(op.get('op'), op.get('op'))}  原件 [{op.get('leftStart')}, {op.get('leftEnd')})  比对件 [{op.get('rightStart')}, {op.get('rightEnd')})"))
    return out, pending


def _docx_bytes(blocks):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_after = Pt(7)
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.2
    doc.styles["Title"].font.size = Pt(24)
    doc.styles["Heading 1"].font.size = Pt(16)
    doc.styles["Heading 2"].font.size = Pt(12)
    for kind, content in blocks:
        if kind == "summary":
            table = doc.add_table(rows=1, cols=2)
            table.autofit = False
            table.columns[0].width, table.columns[1].width = Inches(2.4), Inches(1)
            table.rows[0].cells[0].text, table.rows[0].cells[1].text = "变化类别", "条目数"
            for label, number in content:
                cells = table.add_row().cells
                cells[0].text, cells[1].text = label, str(number)
            for i, row in enumerate(table.rows):
                for cell in row.cells:
                    props = cell._tc.get_or_add_tcPr()
                    borders = OxmlElement("w:tcBorders")
                    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
                        border = OxmlElement("w:" + side)
                        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                            border.set(qn("w:" + key), value)
                        borders.append(border)
                    props.append(borders)
                    margins = OxmlElement("w:tcMar")
                    for side in ("top", "left", "bottom", "right"):
                        margin = OxmlElement("w:" + side)
                        margin.set(qn("w:w"), "100")
                        margin.set(qn("w:type"), "dxa")
                        margins.append(margin)
                    props.append(margins)
                    if i == 0:
                        shading = OxmlElement("w:shd")
                        shading.set(qn("w:fill"), "DCE6F1")
                        props.append(shading)
            doc.add_paragraph()
            continue
        style = {"title": "Title", "h1": "Heading 1", "h2": "Heading 2"}.get(kind)
        p = doc.add_paragraph(content, style)
        if kind == "label":
            p.runs[0].bold = True
            p.paragraph_format.keep_with_next = True
        elif kind == "quote":
            p.paragraph_format.left_indent = Inches(0.15)
        if kind in ("h1", "h2", "title"):
            p.paragraph_format.keep_with_next = True
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def _pdf_bytes(blocks, warnings=None):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    class UnicodeTTFont(TTFont):
        def addObjects(self, doc):
            before=set(doc.idToObject)
            super().addObjects(doc)
            # ReportLab versions using %04X serialize non-BMP destinations as
            # five/six hex digits. PDF ToUnicode requires UTF-16BE surrogates.
            # Repair only CMaps created by this font instance, without changing
            # global ReportLab behavior or font glyph outlines.
            for name in set(doc.idToObject)-before:
                stream=doc.idToObject[name]
                if name.startswith('toUnicodeCMap:') and isinstance(getattr(stream,'content',None),str):
                    stream.content=re.sub(r'(?m)^(<[0-9A-Fa-f]{2}>) <([0-9A-Fa-f]{5,6})>$',
                        lambda m:m[1]+' <'+chr(int(m[2],16)).encode('utf-16-be').hex().upper()+'>',stream.content)
    font = "STSong-Light"
    pdfmetrics.registerFont(UnicodeCIDFont(font))
    symbol_font = None
    candidates = [os.environ.get("DOCUMENT_COMPARE_SYMBOL_FONT"),
                  str(Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts/seguiemj.ttf"),
                  "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"]
    for candidate in candidates:
        if not candidate or not Path(candidate).is_file():
            continue
        try:
            symbol_font = UnicodeTTFont("DocumentCompareSymbols", candidate)
            pdfmetrics.registerFont(symbol_font)
            break
        except Exception:
            continue
    missing = set()
    def text_markup(value, preserve_spaces=False):
        parts = []
        for char in str(value):
            if ord(char) > 0xFFFF:
                if symbol_font is not None and ord(char) in symbol_font.face.charToGlyph:
                    parts.append('<font name="DocumentCompareSymbols">' + char + '</font>')
                else:
                    code = f"U+{ord(char):04X}"
                    missing.add(code)
                    parts.append("[" + code + "]")
            elif char == " " and preserve_spaces:
                parts.append("&#160;")
            elif char == "\t":
                parts.append("&#160;&#160;&#160;&#160;")
            else:
                parts.append(escape(char))
        return "".join(parts) or "&#160;"
    body = ParagraphStyle("Body", fontName=font, fontSize=10.5, leading=16, spaceAfter=7,
                          wordWrap="CJK", splitLongWords=True, textColor=colors.black)
    styles = {"body": body, "quote": ParagraphStyle("Quote", parent=body, leftIndent=10),
              "label": ParagraphStyle("Label", parent=body, keepWithNext=True),
              "title": ParagraphStyle("Title", parent=body, fontSize=24, leading=31, spaceAfter=15, keepWithNext=True),
              "h1": ParagraphStyle("H1", parent=body, fontSize=16, leading=23, spaceBefore=14, spaceAfter=9, keepWithNext=True),
              "h2": ParagraphStyle("H2", parent=body, fontSize=12, leading=18, spaceBefore=10, spaceAfter=7, keepWithNext=True)}
    story = []
    for kind, content in blocks:
        if kind == "summary":
            rows = [[Paragraph("变化类别", body), Paragraph("条目数", body)]] + [[Paragraph(label, body), Paragraph(str(n), body)] for label, n in content]
            table = Table(rows, colWidths=[2.4 * inch, inch], hAlign="LEFT", repeatRows=1)
            table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F1")),
                                       ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
                                       ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                                       ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7)]))
            story += [table, Spacer(1, 10)]
        else:
            # Split long quoted blocks into paragraphs to allow page breaks;
            # XML escaping prevents user text becoming ReportLab markup.
            for line in str(content).split("\n"):
                value = text_markup(line, kind == "quote")
                story.append(Paragraph(value, styles[kind]))
    if missing:
        notice = "字体缺少部分字符，已用码点明确表示 " + "、".join(sorted(missing)) + "。原始字符保留在比对 JSON 中。"
        story.append(Paragraph(notice, body))
        if warnings is not None:
            warnings.append(notice)
    stream = io.BytesIO()
    document = SimpleDocTemplate(stream, title="文档比对独立报告", author="Document Compare", pagesize=(595.28, 841.89),
                                 leftMargin=0.85 * inch, rightMargin=0.85 * inch, topMargin=0.8 * inch, bottomMargin=0.8 * inch)
    def page_number(canvas, doc):
        canvas.setFont(font, 9)
        canvas.drawRightString(595.28 - 0.85 * inch, 0.4 * inch, str(doc.page))
    document.build(story, onFirstPage=page_number, onLaterPages=page_number)
    return stream.getvalue()


def _atomic_write(path, data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(dir=path.parent, prefix=".report-", suffix=path.suffix)
    try:
        with os.fdopen(fd, "wb") as handle:
            handle.write(data)
        os.replace(name, path)
    finally:
        if os.path.exists(name):
            os.unlink(name)


def write_report(result, output_path, format="docx"):
    """Write an independent report plus <output>.manifest.json; return manifest."""
    if format not in ("docx", "pdf"):
        raise ValueError("Report format must be docx or pdf")
    output = Path(output_path)
    if output.suffix.lower() != "." + format:
        raise ValueError("Output extension must match report format")
    blocks, pending = _blocks(result)
    warnings = []
    payload = _docx_bytes(blocks) if format == "docx" else _pdf_bytes(blocks, warnings)
    manifest = {"schemaVersion": "1.0", "kind": "independent_diff_report", "format": format,
                "comparisonId": result.get("comparisonId"), "outputPath": str(output.resolve()),
                "outputSha256": hashlib.sha256(payload).hexdigest(), "alignmentCount": len(result.get("alignments", [])),
                "pendingCount": pending, "includesAllAlignments": True, "includesSourceLayoutIndex": True,
                "semanticJudgmentsAddedByExporter": False, "visualQA": "not_performed_by_writer", "warnings": warnings}
    _atomic_write(output, payload)
    _atomic_write(str(output) + ".manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2).encode("utf-8"))
    return manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("result", help="Comparison JSON")
    parser.add_argument("--output", required=True)
    parser.add_argument("--format", choices=("docx", "pdf"), default="docx")
    args = parser.parse_args()
    try:
        result = json.loads(Path(args.result).read_text(encoding="utf-8-sig"))
        manifest = write_report(result, args.output, args.format)
        print(json.dumps(manifest, ensure_ascii=False, indent=2))
    except (ValueError, OSError) as exc:
        parser.exit(2, f"Report export failed: {exc}\n")


if __name__ == "__main__":
    main()
