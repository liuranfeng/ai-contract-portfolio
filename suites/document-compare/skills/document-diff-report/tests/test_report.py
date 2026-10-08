"""Integration checks for separate report completeness, using synthetic inputs."""
import importlib.util
import sys
import tempfile
import unittest
import io
from pathlib import Path
from docx import Document
from pypdf import PdfReader

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "write_report.py"


class ReportTest(unittest.TestCase):
    def test_report_keeps_every_difference_source_and_pending_reason(self):
        self.assertTrue(SCRIPT.exists(), "Report writer has not been implemented")
        spec = importlib.util.spec_from_file_location("write_report", SCRIPT)
        mod = importlib.util.module_from_spec(spec)
        sys.modules[spec.name] = mod
        spec.loader.exec_module(mod)
        result = {"schemaVersion": "1.0", "comparisonId": "synthetic-fixture", "mode": "semantic",
                  "surface": "reading", "status": "needs_review", "summary": {"replace": 1, "pending": 1},
                  "warnings": ["这是模拟验证数据"], "documents": [
                      {"documentIndex": 0, "name": "Alpha", "parseRunId": "run-A", "layouts": [{"layoutId": "L007", "ordinal": 0, "page": 2, "text": "三年"}]},
                      {"documentIndex": 1, "name": "Beta", "parseRunId": "run-B", "layouts": [{"layoutId": 9, "ordinal": 0, "page": 3, "text": "五年"}]}],
                  "alignments": [{"id": "D0001", "operation": "replace", "alignmentStatus": "aligned", "leftText": "三年", "rightText": "五年",
                      "leftRefs": [{"documentIndex": 0, "layoutId": "L007"}], "rightRefs": [{"documentIndex": 1, "layoutId": 9}],
                      "semantic": {"status": "pending", "reason": "尚待宿主模型判断期限变化"}}]}
        with tempfile.TemporaryDirectory() as td:
            for fmt in ("docx", "pdf"):
                path = Path(td) / ("report." + fmt)
                manifest = mod.write_report(result, path, format=fmt)
                self.assertEqual(manifest["kind"], "independent_diff_report")
                self.assertEqual(manifest["alignmentCount"], 1)
                if fmt == "docx":
                    doc = Document(path)
                    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
                else:
                    text = "\n".join(p.extract_text() for p in PdfReader(path).pages)
                for expected in ("D0001", "L007", "待判断", "三年", "五年", "尚待宿主模型判断期限变化", "run-A", "run-B"):
                    self.assertIn(expected, text)
                self.assertTrue(Path(str(path) + ".manifest.json").exists())

    def test_uncertain_counts_as_review_and_equal_has_no_missing_reason_warning(self):
        self.assertTrue(SCRIPT.exists())
        spec = importlib.util.spec_from_file_location("write_report", SCRIPT)
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        result = {"schemaVersion": "1.0", "mode": "semantic", "documents": [{"name": "甲", "layouts": []}, {"name": "乙", "layouts": []}],
                  "alignments": [{"id": "D1", "operation": "equal", "semantic": {"status": "equal"}},
                                 {"id": "D2", "operation": "replace", "semantic": {"status": "uncertain", "reason": "条件缺失"}}]}
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "report.docx"
            manifest = mod.write_report(result, path)
            text = "\n".join(p.text for p in Document(path).paragraphs)
            self.assertEqual(manifest["pendingCount"], 1)
            self.assertIn("需人工确认", text)
            self.assertNotIn("未提供独立语义理由", text)

    def test_pdf_quote_preserves_significant_spaces(self):
        spec = importlib.util.spec_from_file_location("write_report", SCRIPT)
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        payload = mod._pdf_bytes([("quote", "LEFT    GAP    RIGHT")])
        extracted = PdfReader(io.BytesIO(payload)).pages[0].extract_text().replace("\u00a0", " ")
        self.assertIn("LEFT    GAP    RIGHT", extracted)

    def test_pdf_non_bmp_character_is_visible_or_explicitly_identified(self):
        spec = importlib.util.spec_from_file_location("write_report", SCRIPT)
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        payload = mod._pdf_bytes([("quote", "设备标识 😀")])
        from pypdf import PdfReader
        page = PdfReader(io.BytesIO(payload)).pages[0]
        # CID CJK fonts contain no non-BMP glyph. A dedicated embedded symbol
        # font or an explicit codepoint label is required, never a blank gap.
        fonts = page["/Resources"]["/Font"].get_object()
        has_embedded_symbol = any("DocumentCompareSymbols" in str(font.get_object().get("/BaseFont", "")) or "SegoeUIEmoji" in str(font.get_object().get("/BaseFont", "")) for font in fonts.values())
        self.assertTrue(has_embedded_symbol or "U+1F600" in page.extract_text())
        with self.assertNoLogs('pypdf', level='WARNING'):
            extracted = page.extract_text()
        self.assertTrue('😀' in extracted or 'U+1F600' in extracted, extracted)


if __name__ == "__main__":
    unittest.main()
