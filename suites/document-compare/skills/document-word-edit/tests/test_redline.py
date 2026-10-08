"""Synthetic fixtures, not OCR output or a real client contract.

These tests catch loss of source formatting, incorrect revision projections,
ambiguous targeting, hidden omissions and accidental source overwrites.
"""
import copy
import importlib.util
import json
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree as ET

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "write_redline.py"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def load_writer():
    if not SCRIPT.exists():
        raise AssertionError("Word writer has not been implemented")
    spec = importlib.util.spec_from_file_location("write_redline", SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = mod
    spec.loader.exec_module(mod)
    return mod


def fixture_result(left, right, pairs):
    docs = [{"documentIndex": d, "name": name, "parseRunId": "synthetic-fixture",
             "layouts": [{"layoutId": f"{d}-{i}", "ordinal": i, "page": 1,
                          "type": "text", "text": t, "readingText": t}
                         for i, t in enumerate(texts)]}
            for d, name, texts in [(0, "模拟原件", left), (1, "模拟比对件", right)]]
    rows = []
    for n, (li, ri) in enumerate(pairs):
        lt, rt = "\n".join(left[i] for i in li), "\n".join(right[i] for i in ri)
        op = "insert" if not li else "delete" if not ri else "equal" if lt == rt else "replace"
        rows.append({"id": f"D{n+1:04}", "left": li, "right": ri, "operation": op,
                     "alignmentStatus": "aligned", "leftText": lt, "rightText": rt,
                     "leftRefs": [{"documentIndex": 0, "layoutId": f"0-{i}"} for i in li],
                     "rightRefs": [{"documentIndex": 1, "layoutId": f"1-{i}"} for i in ri],
                     "semantic": {"status": "not_requested"}})
    return {"schemaVersion": "1.0", "mode": "character", "surface": "reading",
            "comparisonId": "fixture", "status": "complete", "documents": docs,
            "alignments": rows, "summary": {"pending": 0}, "warnings": []}


def projected_paragraphs(path, accept):
    """Independent OOXML projection, honoring paragraph marks and run revisions."""
    with ZipFile(path) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    out = []
    for p in root.iter(W + "p"):
        mark = p.find(W + "pPr/" + W + "rPr")
        if mark is not None and mark.find(W + ("del" if accept else "ins")) is not None:
            continue
        chunks = []
        for el in p.iter():
            ancestors = {x.tag for x in el.iterancestors()}
            if (accept and W + "del" in ancestors) or (not accept and W + "ins" in ancestors):
                continue
            if el.tag in {W + "t", W + "delText"}:
                chunks.append(el.text or "")
            elif el.tag == W + "tab":
                chunks.append("\t")
            elif el.tag in {W + "br", W + "cr"}:
                chunks.append("\n")
        out.append("".join(chunks))
    return out


class RedlineTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.src = self.root / "source.docx"
        self.out = self.root / "redline.docx"
        self.writer = load_writer()

    def tearDown(self):
        self.tmp.cleanup()

    def make_source(self):
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("期限").bold = True
        p.add_run("三年").italic = True
        doc.add_paragraph("应删除条款")
        doc.add_paragraph("最后条款")
        doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "费用100元"
        doc.tables[0].cell(0, 1).text = "每月支付"
        doc.sections[0].header.paragraphs[0].text = "原始页眉"
        doc.sections[0].footer.paragraphs[0].text = "原始页脚"
        doc.save(self.src)

    def test_real_revisions_roundtrip_and_package_preservation(self):
        self.make_source()
        baseline = self.src.read_bytes()
        result = fixture_result(["期限三年", "应删除条款", "最后条款", "费用100元", "每月支付"],
                                ["期限五年", "新增条款", "最后条款", "费用120元", "每月支付"],
                                [([0], [0]), ([1], []), ([], [1]), ([2], [2]), ([3], [3]), ([4], [4])])
        for _ in range(2):
            manifest = self.writer.write_redline(result, self.src, self.out)
            self.assertEqual(projected_paragraphs(self.out, False),
                             ["期限三年", "应删除条款", "最后条款", "费用100元", "每月支付"])
            self.assertEqual(projected_paragraphs(self.out, True),
                             ["期限五年", "新增条款", "最后条款", "费用120元", "每月支付"])
            self.assertTrue(manifest["validation"]["rejectEqualsSource"])
            self.assertTrue(manifest["validation"]["acceptEqualsTarget"])
        self.assertEqual(baseline, self.src.read_bytes())
        with ZipFile(self.src) as src, ZipFile(self.out) as out:
            for name in src.namelist():
                if name not in {"word/document.xml", "word/settings.xml"}:
                    self.assertEqual(src.read(name), out.read(name), name)
            root = ET.fromstring(out.read("word/document.xml"))
            self.assertTrue(root.findall(".//" + W + "ins"))
            self.assertTrue(root.findall(".//" + W + "delText"))
            unchanged = root.find(".//" + W + "r")
            self.assertIsNotNone(unchanged.find(W + "rPr/" + W + "b"))

    def test_ambiguous_text_requires_explicit_mapping(self):
        doc = Document()
        doc.add_paragraph("重复")
        doc.add_paragraph("重复")
        doc.save(self.src)
        result = fixture_result(["重复", "重复"], ["重复", "已改"], [([0], [0]), ([1], [1])])
        with self.assertRaisesRegex(ValueError, "ambiguous|mapping"):
            self.writer.write_redline(result, self.src, self.out)
        mapping = {"layoutMap": [{"layoutId": "0-0", "paragraphIndex": 0},
                                  {"layoutId": "0-1", "paragraphIndex": 1}]}
        self.writer.write_redline(result, self.src, self.out, mapping)
        self.assertEqual(projected_paragraphs(self.out, True), ["重复", "已改"])

    def test_existing_revisions_and_source_overwrite_rejected(self):
        doc = Document()
        doc.add_paragraph("原文")
        doc.save(self.src)
        result = fixture_result(["原文"], ["新文"], [([0], [0])])
        with self.assertRaisesRegex(ValueError, "overwrite|source"):
            self.writer.write_redline(result, self.src, self.src)
        self.writer.write_redline(result, self.src, self.out)
        with self.assertRaisesRegex(ValueError, "existing revisions"):
            self.writer.write_redline(result, self.out, self.root / "second.docx")

    def test_unmapped_source_text_cannot_be_silently_omitted(self):
        self.make_source()
        result = fixture_result(["期限三年"], ["期限五年"], [([0], [0])])
        with self.assertRaisesRegex(ValueError, "coverage|unmapped"):
            self.writer.write_redline(result, self.src, self.out)
        self.assertFalse(self.out.exists())

    def test_mismatched_mapping_and_unresolved_alignment_rejected(self):
        self.make_source()
        result = fixture_result(["假原文"], ["新文"], [([0], [0])])
        with self.assertRaises(ValueError):
            self.writer.write_redline(result, self.src, self.out,
                {"layoutMap": [{"layoutId": "0-0", "paragraphIndex": 0}]})
        result["alignments"][0]["alignmentStatus"] = "unresolved"
        with self.assertRaisesRegex(ValueError, "unresolved"):
            self.writer.write_redline(result, self.src, self.out)

    def test_reconstruction_is_explicit_and_labelled(self):
        result = fixture_result(["第一条", "删除条款"], ["第一条修订"], [([0], [0]), ([1], [])])
        with self.assertRaises(ValueError):
            self.writer.write_redline(result, None, self.out)
        manifest = self.writer.write_redline(result, None, self.out, reconstructed=True)
        self.assertEqual(manifest["mode"], "reconstructed")
        self.assertIn("重建", projected_paragraphs(self.out, True)[0])
        self.assertEqual(projected_paragraphs(self.out, True)[1:], ["第一条修订"])

    def test_table_cell_mapping_and_complex_run_refusal(self):
        doc = Document()
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "旧值"
        doc.save(self.src)
        result = fixture_result(["旧值"], ["新值"], [([0], [0])])
        self.writer.write_redline(result, self.src, self.out, {"layoutMap": [
            {"layoutId": "0-0", "tableIndex": 0, "rowIndex": 0, "cellIndex": 0, "cellParagraphIndex": 0}]})
        self.assertEqual(projected_paragraphs(self.out, True), ["新值"])
        doc = Document()
        p = doc.add_paragraph("旧值")
        fld = ET.SubElement(p.runs[0]._r, W + "fldChar")
        fld.set(W + "fldCharType", "begin")
        doc.save(self.src)
        with self.assertRaisesRegex(ValueError, "complex|unsupported"):
            self.writer.write_redline(result, self.src, self.out)

    def test_whole_table_layout_updates_cells_and_rejects_shape_change(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        for row, values in zip(table.rows, [["项目", "金额"], ["维护费", "100元"]]):
            for cell, value in zip(row.cells, values):
                cell.text = value
        doc.save(self.src)
        result = fixture_result(["项目\t金额\n维护费\t100元"], ["项目\t金额\n维护费\t120元"], [([0], [0])])
        for mapping in (None, {"layoutMap": [{"layoutId": "0-0", "tableIndex": 0}]}):
            self.writer.write_redline(result, self.src, self.out, mapping)
            self.assertEqual(projected_paragraphs(self.out, False), ["项目", "金额", "维护费", "100元"])
            self.assertEqual(projected_paragraphs(self.out, True), ["项目", "金额", "维护费", "120元"])
        result = fixture_result(["项目\t金额\n维护费\t100元"], ["项目\t金额\n维护费\t120元\n新增费\t1元"], [([0], [0])])
        with self.assertRaisesRegex(ValueError, "table.*structure|table.*shape"):
            self.writer.write_redline(result, self.src, self.out)

    def test_insert_after_deleted_last_paragraph_does_not_inherit_deletion(self):
        doc = Document()
        doc.add_paragraph("原条款")
        doc.save(self.src)
        result = fixture_result(["原条款"], ["新条款"], [([0], []), ([], [0])])
        self.writer.write_redline(result, self.src, self.out)
        self.assertEqual(projected_paragraphs(self.out, False), ["原条款"])
        self.assertEqual(projected_paragraphs(self.out, True), ["新条款"])

    def test_native_table_insertion_is_not_silently_flattened(self):
        doc = Document()
        doc.add_paragraph("原条款")
        doc.save(self.src)
        result = fixture_result(["原条款"], ["原条款", "项目\t金额\n服务\t100"], [([0], [0]), ([], [1])])
        result["documents"][1]["layouts"][1]["type"] = "table"
        with self.assertRaisesRegex(ValueError, "table.*insert|table.*structure"):
            self.writer.write_redline(result, self.src, self.out)

    def test_consecutive_tail_insertions_match_actual_target_order(self):
        doc = Document()
        doc.add_paragraph("A")
        doc.save(self.src)
        result = fixture_result(["A"], ["A", "B", "C"], [([0], [0]), ([], [1]), ([], [2])])
        self.writer.write_redline(result, self.src, self.out)
        self.assertEqual(projected_paragraphs(self.out, True), ["A", "B", "C"])


if __name__ == "__main__":
    unittest.main()
