#!/usr/bin/env python3
"""Conservative, source-package-preserving Word tracked changes.

No model, OCR or network calls. Input is the document comparison JSON contract.
Ambiguous mapping, existing revisions and unsupported edited OOXML fail closed.
"""
import argparse
import copy
import difflib
import hashlib
import io
import json
import os
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZipFile

from lxml import etree as ET

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
XML = "{http://www.w3.org/XML/1998/namespace}"
NS = {"w": W[1:-1]}
MAIN = "word/document.xml"
RECONSTRUCTION_TITLE = "文档比对修订稿 由解析文本重建"
PARSER = ET.XMLParser(resolve_entities=False, no_network=True)


def _read_xml(data):
    return ET.fromstring(data, parser=PARSER)


def _xml_bytes(root):
    return ET.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _text(node, accept=None):
    chunks = []
    for el in node.iter():
        ancestors = {a.tag for a in el.iterancestors()}
        if (accept is True and W + "del" in ancestors) or (accept is False and W + "ins" in ancestors):
            continue
        if el.tag in (W + "t", W + "delText"):
            chunks.append(el.text or "")
        elif el.tag == W + "tab":
            chunks.append("\t")
        elif el.tag in (W + "br", W + "cr"):
            chunks.append("\n")
    return "".join(chunks)


def _projection(root, accept):
    result = []
    for p in root.iter(W + "p"):
        mark = p.find(W + "pPr/" + W + "rPr/" + W + ("del" if accept else "ins"))
        if mark is None:
            result.append(_text(p, accept))
    return result


def _surface(layout, surface):
    value = layout.get("readingText", layout.get("text", "")) if surface == "reading" else layout.get("text", "")
    if not isinstance(value, str):
        raise ValueError("layout text must be a string")
    return value


def _validate_result(result):
    if str(result.get("schemaVersion")) != "1.0" or len(result.get("documents", [])) != 2:
        raise ValueError("Expected schemaVersion 1.0 and exactly two documents")
    docs = result["documents"]
    seen = [[], []]
    for item in result.get("alignments", []):
        if item.get("alignmentStatus") == "unresolved":
            raise ValueError(f"unresolved alignment {item.get('id')}; resolve mapping before Word output")
        for side, key in enumerate(("left", "right")):
            indices = item.get(key, [])
            if any(type(i) is not int or i < 0 or i >= len(docs[side]["layouts"]) for i in indices):
                raise ValueError("Invalid layout index")
            seen[side].extend(indices)
            want = "\n".join(_surface(docs[side]["layouts"][i], result.get("surface", "reading")) for i in indices)
            if item.get(key + "Text") != want:
                raise ValueError(f"{item.get('id')}: {key}Text differs from indexed layouts")
        if not item.get("left") and not item.get("right"):
            raise ValueError("Alignment cannot have two empty sides")
    for side in (0, 1):
        if seen[side] != list(range(len(docs[side]["layouts"]))):
            raise ValueError("Alignment coverage must be complete, unique and monotonic")


def _plain_editable(p):
    for c in p:
        if c.tag not in (W + "pPr", W + "r"):
            raise ValueError("unsupported complex paragraph; fields, hyperlinks, bookmarks and controls require explicit expert OOXML editing")
        if c.tag == W + "r":
            for el in c:
                if el.tag not in (W + "rPr", W + "t", W + "tab", W + "br", W + "cr"):
                    raise ValueError("unsupported complex run including fields, drawings or embedded objects")
                if el.tag == W + "br" and el.get(W + "type", "textWrapping") != "textWrapping":
                    raise ValueError("unsupported complex run containing page or column break")
    if p.find(".//" + W + "vanish") is not None:
        raise ValueError("unsupported complex hidden text in edited paragraph")


def _text_run(text, template=None, deleted=False):
    run = copy.deepcopy(template) if template is not None else ET.Element(W + "r")
    for child in list(run):
        if child.tag != W + "rPr":
            run.remove(child)
    start = 0
    for i, ch in enumerate(text):
        if ch not in "\t\n":
            continue
        if i > start:
            el = ET.SubElement(run, W + ("delText" if deleted else "t"))
            el.set(XML + "space", "preserve")
            el.text = text[start:i]
        ET.SubElement(run, W + ("tab" if ch == "\t" else "br"))
        start = i + 1
    if start < len(text):
        el = ET.SubElement(run, W + ("delText" if deleted else "t"))
        el.set(XML + "space", "preserve")
        el.text = text[start:]
    return run


class _Revisions:
    def __init__(self, first_id):
        self.next_id = first_id
        self.timestamp = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")

    def tag(self, kind):
        el = ET.Element(W + kind)
        el.set(W + "id", str(self.next_id))
        el.set(W + "author", "Document Compare")
        el.set(W + "date", self.timestamp)
        self.next_id += 1
        return el

    def paragraph_mark(self, p, kind):
        props = p.find(W + "pPr")
        if props is None:
            props = ET.Element(W + "pPr")
            p.insert(0, props)
        rp = props.find(W + "rPr")
        if rp is None:
            # rPr belongs before sectPr and pPrChange in the pPr sequence.
            rp = ET.Element(W + "rPr")
            section = props.find(W + "sectPr")
            props.insert(props.index(section) if section is not None else len(props), rp)
        rp.append(self.tag(kind))

    def replace(self, p, target):
        _plain_editable(p)
        old = _text(p)
        runs, cursor = [], 0
        for r in p.findall(W + "r"):
            value = _text(r)
            runs.append((cursor, cursor + len(value), r, value))
            cursor += len(value)
        def pieces(begin, end, deleted=False):
            return [_text_run(value[max(begin - a, 0):min(end - a, b - a)], r, deleted)
                    for a, b, r, value in runs if b > begin and a < end]
        for c in list(p):
            if c.tag != W + "pPr":
                p.remove(c)
        for kind, a, b, c, d in difflib.SequenceMatcher(None, old, target, autojunk=False).get_opcodes():
            if kind == "equal":
                p.extend(pieces(a, b))
                continue
            if kind in ("delete", "replace"):
                deleted = self.tag("del")
                deleted.extend(pieces(a, b, True))
                p.append(deleted)
            if kind in ("insert", "replace"):
                template = next((r for x, y, r, _ in runs if x <= a < y), runs[-1][2] if runs else None)
                inserted = self.tag("ins")
                inserted.append(_text_run(target[c:d], template))
                p.append(inserted)

    def delete_paragraph(self, p):
        if p.find(W + "pPr/" + W + "sectPr") is not None:
            raise ValueError("unsupported deletion of a paragraph carrying a section break")
        if p.getparent().tag == W + "tc":
            raise ValueError("unsupported structural deletion inside a table cell; provide a cell-content replacement or expert table edit")
        self.replace(p, "")
        self.paragraph_mark(p, "del")

    def insert_paragraph(self, anchor, target, before=False):
        p = ET.Element(W + "p")
        props = anchor.find(W + "pPr")
        if props is not None:
            props = copy.deepcopy(props)
            for section in list(props.findall(W + "sectPr")):
                props.remove(section)
            # The anchor may already be scheduled for deletion. Only its
            # original formatting is inherited, never its revision marks.
            for mark in list(props.findall(".//" + W + "ins")) + list(props.findall(".//" + W + "del")):
                mark.getparent().remove(mark)
            p.append(props)
        self.paragraph_mark(p, "ins")
        inserted = self.tag("ins")
        inserted.append(_text_run(target, anchor.find(W + "r")))
        p.append(inserted)
        if before:
            anchor.addprevious(p)
        else:
            anchor.addnext(p)
        return p


def inspect_docx(source_docx):
    """Return stable paragraphIndex values for explicit mapping, without editing."""
    rows, tables = [], []
    with ZipFile(source_docx) as package:
        for name in package.namelist():
            if name.startswith("word/") and name.endswith(".xml"):
                root = _read_xml(package.read(name))
                ps = list(root.iter(W + "p"))
                for i, p in enumerate(root.iter(W + "p")):
                    rows.append({"part": name, "paragraphIndex": i, "text": _text(p),
                                 "inTable": any(a.tag == W + "tc" for a in p.iterancestors())})
                for ti, table in enumerate(root.iter(W + "tbl")):
                    cells = []
                    for ri, row in enumerate(table.findall(W + "tr")):
                        for ci, cell in enumerate(row.findall(W + "tc")):
                            for pi, p in enumerate(cell.findall(W + "p")):
                                cells.append({"rowIndex": ri, "cellIndex": ci, "cellParagraphIndex": pi,
                                              "paragraphIndex": ps.index(p), "text": _text(p)})
                    tables.append({"part": name, "tableIndex": ti, "cells": cells})
    return {"sourceSha256": hashlib.sha256(Path(source_docx).read_bytes()).hexdigest(), "paragraphs": rows, "tables": tables}


def _table_cells(table):
    rows = table.findall(W + "tr")
    shape, selected, texts = [], [], []
    for row in rows:
        cells = row.findall(W + "tc")
        shape.append(len(cells))
        values = []
        for cell in cells:
            ps = cell.findall(W + "p")
            if len(ps) != 1 or cell.find(W + "tbl") is not None:
                raise ValueError("unsupported table layout with nested tables or multiple paragraphs per cell; split into paragraph layouts and map cellParagraphIndex explicitly")
            if cell.find(W + "tcPr/" + W + "gridSpan") is not None or cell.find(W + "tcPr/" + W + "vMerge") is not None:
                raise ValueError("unsupported merged-cell table layout; map individual physical cell paragraphs explicitly")
            value = _text(ps[0])
            if "\t" in value or "\n" in value:
                raise ValueError("unsupported table layout with internal cell tabs/newlines; map individual cell paragraphs explicitly")
            selected.append(ps[0])
            values.append(value)
        texts.append("\t".join(values))
    return shape, selected, "\n".join(texts)


def _resolve_mapping(result, roots, mapping, reconstructed=False):
    paragraphs = {part: list(root.iter(W + "p")) for part, root in roots.items()}
    layouts = result["documents"][0]["layouts"]
    entries = (mapping or {}).get("layoutMap", [])
    by_id = {}
    for entry in entries:
        key = (type(entry["layoutId"]).__name__, str(entry["layoutId"]))
        if key in by_id:
            raise ValueError("Duplicate layoutId in mapping")
        by_id[key] = entry
    resolved, used, audit, table_shapes = {}, set(), [], {}
    for i, layout in enumerate(layouts):
        expected = _surface(layout, result.get("surface", "reading"))
        key = (type(layout["layoutId"]).__name__, str(layout["layoutId"]))
        entry = by_id.get(key)
        part = entry.get("part", MAIN) if entry else MAIN
        if part not in roots:
            raise ValueError(f"mapping references missing XML part {part}")
        ps = paragraphs[part]
        table_shape = None
        if entry is None:
            candidates = [j for j, p in enumerate(ps) if _text(p) == expected]
            tables = []
            for table in roots[part].iter(W + "tbl"):
                try:
                    shape, selected, value = _table_cells(table)
                except ValueError:
                    continue
                if value == expected:
                    tables.append((shape, [ps.index(p) for p in selected]))
            # A one-cell table is also one paragraph; those identical targets
            # represent a single location, not a duplicate-text ambiguity.
            if len(candidates) == 1 and len(tables) == 1 and tables[0][1] == candidates:
                tables = []
            if len(candidates) + len(tables) != 1:
                raise ValueError(f"ambiguous or missing mapping for layoutId {layout['layoutId']!r}; inspect source and supply layoutMap")
            if tables:
                table_shape, indices = tables[0]
            else:
                indices = candidates
        elif "tableIndex" in entry:
            try:
                table = list(roots[part].iter(W + "tbl"))[entry["tableIndex"]]
                if "rowIndex" not in entry:
                    table_shape, selected, _ = _table_cells(table)
                    indices = [ps.index(p) for p in selected]
                else:
                    cell = table.findall(W + "tr")[entry["rowIndex"]].findall(W + "tc")[entry["cellIndex"]]
                    p = cell.findall(W + "p")[entry.get("cellParagraphIndex", 0)]
                    indices = [ps.index(p)]
            except (IndexError, KeyError, TypeError, ValueError) as exc:
                raise ValueError("Invalid physical table row/cell/paragraph mapping") from exc
        else:
            indices = entry.get("paragraphIndices", [entry.get("paragraphIndex")])
        if not indices or any(type(j) is not int or j < 0 or j >= len(ps) for j in indices):
            raise ValueError("Invalid paragraphIndex in mapping")
        if indices != sorted(set(indices)):
            raise ValueError("mapping paragraphIndices must be unique and ordered")
        if any((part, j) in used for j in indices):
            raise ValueError("mapping duplicates an OOXML paragraph")
        selected = [ps[j] for j in indices]
        if table_shape is not None:
            values, start = [], 0
            for width in table_shape:
                values.append("\t".join(_text(p) for p in selected[start:start + width]))
                start += width
            mapped_text = "\n".join(values)
            table_shapes[i] = table_shape
        else:
            mapped_text = "\n".join(_text(p) for p in selected)
        if mapped_text != expected:
            raise ValueError(f"mapping text does not match source layoutId {layout['layoutId']!r}")
        used.update((part, j) for j in indices)
        resolved[i] = [(part, p) for p in selected]
        audit.append({"layoutId": layout["layoutId"], "documentIndex": 0, "part": part, "paragraphIndices": indices,
                      **({"tableShape": table_shape} if table_shape else {})})
    # Source text cannot disappear because OCR omitted a block. Empty formatting
    # paragraphs remain untouched, as do non-targeted headers/footnotes (reported).
    covered_parts = {MAIN} | {part for part, _ in used}
    for part in covered_parts:
        for j, p in enumerate(paragraphs[part]):
            if _text(p) and (part, j) not in used and not (reconstructed and part == MAIN and j == 0):
                raise ValueError(f"unmapped source text: coverage missing {part} paragraphIndex {j}")
    linear = [(part, paragraphs[part].index(p)) for i in range(len(layouts)) for part, p in resolved[i]]
    for part in covered_parts:
        order = [j for name, j in linear if name == part]
        if order != sorted(order):
            raise ValueError("mapping does not follow source reading order")
    unknown = set(by_id) - {(type(l["layoutId"]).__name__, str(l["layoutId"])) for l in layouts}
    if unknown:
        raise ValueError("mapping contains unknown layoutIds")
    return resolved, audit, paragraphs, table_shapes


def _atomic_bytes(path, data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix=".compare-", suffix=path.suffix, dir=path.parent)
    try:
        with os.fdopen(fd, "wb") as handle:
            handle.write(data)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def write_redline(result, source_docx, output_path, mapping=None, reconstructed=False):
    """Write true tracked DOCX and its manifest, or raise ValueError before output.

    mapping = {layoutMap: [{layoutId, part?, paragraphIndex|paragraphIndices|tableIndex,
                           rowIndex?, cellIndex?, cellParagraphIndex?}],
               insertions: {alignmentId: {part?, beforeParagraphIndex|afterParagraphIndex}}}
    Indices are zero-based physical OOXML indices; inspect_docx exposes them.
    """
    _validate_result(result)
    output = Path(output_path)
    if output.suffix.lower() != ".docx":
        raise ValueError("Word output must have .docx extension")
    source = Path(source_docx) if source_docx else None
    if source and source.resolve() == output.resolve():
        raise ValueError("Refusing to overwrite source DOCX")
    warnings = list(result.get("warnings", []))
    if reconstructed:
        from docx import Document
        doc = Document()
        doc.add_paragraph(RECONSTRUCTION_TITLE, "Title")
        for layout in result["documents"][0]["layouts"]:
            doc.add_paragraph(_surface(layout, result.get("surface", "reading")))
        if not result["documents"][0]["layouts"]:
            doc.add_paragraph("")
        stream = io.BytesIO()
        doc.save(stream)
        baseline_bytes = stream.getvalue()
        mapping = {"layoutMap": [{"layoutId": layout["layoutId"], "paragraphIndex": i + 1}
                                  for i, layout in enumerate(result["documents"][0]["layouts"])]}
        warnings.append("由解析文本重建，未保留原 PDF 或 Word 的原始版式；标题为永久说明。")
    else:
        if source is None or source.suffix.lower() != ".docx":
            raise ValueError("An original source DOCX is required; select reconstructed=True explicitly for PDF or parsed text")
        baseline_bytes = source.read_bytes()
        expected_hash = result["documents"][0].get("sourceSha256")
        if expected_hash and expected_hash != hashlib.sha256(baseline_bytes).hexdigest():
            raise ValueError("source DOCX SHA256 does not match comparison input")
    with ZipFile(io.BytesIO(baseline_bytes)) as package:
        infos = package.infolist()
        blobs = {info.filename: package.read(info) for info in infos}
    if len(blobs) != len(infos):
        raise ValueError("Unsupported DOCX with duplicate ZIP member names")
    if any(name.startswith("_xmlsignatures/") for name in blobs):
        raise ValueError("Digitally signed source DOCX cannot be edited without invalidating its signature")
    roots = {name: _read_xml(data) for name, data in blobs.items()
             if name.startswith("word/") and name.endswith(".xml")}
    if MAIN not in roots or "word/settings.xml" not in roots:
        raise ValueError("Unsupported DOCX missing document.xml or settings.xml")
    revision_names = {"ins", "del", "moveFrom", "moveTo", "moveFromRangeStart", "moveToRangeStart", "cellIns", "cellDel", "cellMerge", "numberingChange"}
    ids = []
    for name, root in roots.items():
        for el in root.iter():
            local = ET.QName(el).localname
            if el.tag.startswith(W) and (local in revision_names or local.endswith("PrChange")):
                raise ValueError(f"existing revisions in {name}; choose and save an accepted/rejected baseline first")
            value = el.get(W + "id")
            if value and value.isdigit():
                ids.append(int(value))
    resolved, audit, original_paragraphs, table_shapes = _resolve_mapping(result, roots, mapping, reconstructed)
    snapshots = {part: [_text(p) for p in ps] for part, ps in original_paragraphs.items()}
    # Expected projections are a separate text plan, keyed by original element.
    expected = {part: [(p, _text(p)) for p in ps] for part, ps in original_paragraphs.items()}
    revisions = _Revisions(max(ids, default=0) + 1)
    changed_parts = set()
    emitted_groups, insertion_tails = [], {}
    inserted_count = deleted_count = replaced_count = 0
    for n, item in enumerate(result["alignments"]):
        left = [pair for i in item["left"] for pair in resolved[i]]
        targets = [_surface(result["documents"][1]["layouts"][i], result.get("surface", "reading")) for i in item["right"]]
        if item["leftText"] == item["rightText"] and left:
            emitted_groups.append((item, [(resolved[i], table_shapes.get(i)) for i in item["left"]]))
            continue
        if not reconstructed and any(result["documents"][1]["layouts"][i].get("type") == "table" for i in item["right"]) and not any(i in table_shapes for i in item["left"]):
            raise ValueError(f"{item['id']}: unsupported table insertion or table structure creation in original DOCX mode")
        output_nodes = []
        shape = None
        if any(i in table_shapes for i in item["left"]):
            if len(item["left"]) != 1 or len(targets) != 1:
                raise ValueError(f"{item['id']}: unsupported table structure change or table grouped with other layouts")
            rows = [row.split("\t") for row in targets[0].split("\n")]
            if [len(row) for row in rows] != table_shapes[item["left"][0]]:
                raise ValueError(f"{item['id']}: unsupported table shape change; preserve rows and cells for tracked cell edits")
            shape = table_shapes[item["left"][0]]
            targets = [cell for row in rows for cell in row]
        if left:
            parts = {part for part, p in left}
            if len(parts) != 1:
                raise ValueError(f"{item['id']}: unsupported replacement across XML parts")
            part = left[0][0]
            if len(left) != len(targets) and len({p.getparent() for _, p in left}) > 1:
                raise ValueError(f"{item['id']}: unsupported structural change across table cells or containers")
            for j, (_, p) in enumerate(left):
                plan = expected[part]
                at = next(k for k, (node, _) in enumerate(plan) if node is p)
                if j < len(targets):
                    if _text(p) != targets[j]:
                        revisions.replace(p, targets[j])
                        replaced_count += 1
                    plan[at] = (p, targets[j])
                    output_nodes.append((part, p))
                else:
                    revisions.delete_paragraph(p)
                    plan.pop(at)
                    deleted_count += 1
            anchor = left[-1][1]
            for target in targets[len(left):]:
                p = revisions.insert_paragraph(anchor, target)
                plan = expected[part]
                at = next(k for k, (node, _) in enumerate(plan) if node is anchor)
                plan.insert(at + 1, (p, target))
                anchor = p
                output_nodes.append((part, p))
                inserted_count += 1
        else:
            explicit = (mapping or {}).get("insertions", {}).get(item["id"])
            before = False
            if explicit:
                part = explicit.get("part", MAIN)
                before = "beforeParagraphIndex" in explicit
                key = "beforeParagraphIndex" if before else "afterParagraphIndex"
                index = explicit.get(key)
                if part not in original_paragraphs or type(index) is not int or not 0 <= index < len(original_paragraphs[part]):
                    raise ValueError("Invalid explicit insertion anchor")
                anchor = original_paragraphs[part][index]
            else:
                previous = next((resolved[a["left"][-1]][-1] for a in reversed(result["alignments"][:n]) if a["left"]), None)
                following = next((resolved[a["left"][0]][0] for a in result["alignments"][n+1:] if a["left"]), None)
                if previous and following and (previous[0] != following[0] or previous[1].getparent() is not following[1].getparent()):
                    raise ValueError(f"{item['id']}: ambiguous insertion across containers; provide mapping.insertions")
                if following:
                    part, anchor = following
                    before = True
                elif previous:
                    part, anchor = previous
                elif reconstructed:
                    part, anchor = MAIN, original_paragraphs[MAIN][0]
                else:
                    raise ValueError("Insertion into empty source requires explicit mapping.insertions")
            gap = (part, id(anchor), before)
            if gap in insertion_tails:
                anchor, before = insertion_tails[gap], False
            for target in targets:
                p = revisions.insert_paragraph(anchor, target, before)
                plan = expected[part]
                # Deleted anchor can still exist in the XML. Locate its nearest
                # surviving neighbor in physical order for the expected plan.
                position = list(roots[part].iter(W + "p")).index(p)
                physical = list(roots[part].iter(W + "p"))
                at = sum(physical.index(node) < position for node, _ in plan)
                plan.insert(at, (p, target))
                anchor, before = p, False
                output_nodes.append((part, p))
                inserted_count += 1
            insertion_tails[gap] = anchor
        changed_parts.add(part)
        emitted_groups.append((item, [(output_nodes, shape)]))
    # Independently reconstruct each right-side alignment from actual OOXML
    # nodes, then check their physical ordering against the right layout order.
    # This catches e.g. reversed tail insertions even if a mutation plan shares
    # the same bug. Unmapped formatting paragraphs are excluded only here.
    positions = {part: {p: i for i, p in enumerate(root.iter(W + "p"))} for part, root in roots.items()}
    last_position, actual_groups = {}, []
    for item, segments in emitted_groups:
        segment_texts = []
        for pairs, shape in segments:
            values = []
            for part, p in pairs:
                pos = positions[part][p]
                if pos <= last_position.get(part, -1):
                    raise ValueError("Accepting revisions does not follow target layout reading order")
                last_position[part] = pos
                values.append(_text(p, True))
            if shape:
                rows, cursor = [], 0
                for width in shape:
                    rows.append("\t".join(values[cursor:cursor + width]))
                    cursor += width
                segment_texts.append("\n".join(rows))
            else:
                segment_texts.append("\n".join(values))
        actual = "\n".join(segment_texts)
        if actual != item["rightText"]:
            raise ValueError(f"{item['id']}: accepted OOXML text does not equal target alignment text")
        if item["right"]:
            actual_groups.append(actual)
    independent_target = "\n".join(_surface(layout, result.get("surface", "reading")) for layout in result["documents"][1]["layouts"])
    if "\n".join(actual_groups) != independent_target:
        raise ValueError("Accepted OOXML projection differs from target document layouts")
    for part, root in roots.items():
        if _projection(root, False) != snapshots[part]:
            raise ValueError(f"Internal validation failed: rejecting revisions differs from source in {part}")
        if _projection(root, True) != [text for _, text in expected[part]]:
            raise ValueError(f"Internal validation failed: accepting revisions differs from target plan in {part}")
    settings = roots["word/settings.xml"]
    if settings.find(W + "trackRevisions") is None:
        settings.append(ET.Element(W + "trackRevisions"))
    changed_parts.add("word/settings.xml")
    for part in changed_parts:
        blobs[part] = _xml_bytes(roots[part])
    unmapped_parts = [part for part, ps in original_paragraphs.items()
                      if part != MAIN and any(_text(p) for p in ps) and part not in {a["part"] for a in audit}]
    if unmapped_parts:
        warnings.append("保留但未纳入文字比对的 OOXML 部件: " + ", ".join(unmapped_parts))
    if any(root.find(".//" + W + "drawing") is not None or root.find(".//" + W + "pict") is not None for root in roots.values()):
        warnings.append("原件图形原样保留；此文字修订不验证图片内容是否相同。")
    stream = io.BytesIO()
    with ZipFile(stream, "w") as package:
        for info in infos:
            package.writestr(info, blobs[info.filename])
    generated = stream.getvalue()
    # Re-open the serialized package before publishing; a malformed ZIP is never
    # reported as successful merely because an in-memory projection passed.
    with ZipFile(io.BytesIO(generated)) as check:
        if check.testzip() is not None:
            raise ValueError("Output ZIP CRC validation failed")
        for part in changed_parts - {"word/settings.xml"}:
            root = _read_xml(check.read(part))
            if _projection(root, False) != snapshots[part] or _projection(root, True) != [t for _, t in expected[part]]:
                raise ValueError("Serialized revision validation failed")
    manifest = {"schemaVersion": "1.0", "kind": "tracked_word", "mode": "reconstructed" if reconstructed else "original_docx",
                "comparisonId": result.get("comparisonId"), "sourcePath": str(source.resolve()) if source else None,
                "sourceSha256": hashlib.sha256(source.read_bytes()).hexdigest() if source else None,
                "outputPath": str(output.resolve()), "outputSha256": hashlib.sha256(generated).hexdigest(),
                "mapping": audit, "changedParts": sorted(changed_parts), "preservedUncomparedParts": unmapped_parts,
                "changes": {"insertParagraphs": inserted_count, "deleteParagraphs": deleted_count, "replaceParagraphs": replaced_count},
                "validation": {"rejectEqualsSource": True, "acceptEqualsTarget": True, "allPackagePartsChecked": True,
                               "targetLayoutsIndependentlyChecked": True,
                               "scope": "Full paragraph text in every OOXML part, including retained empty paragraphs and uncompared parts",
                               "visualQA": "not_performed_by_writer"},
                "warnings": warnings}
    _atomic_bytes(output, generated)
    _atomic_bytes(str(output) + ".manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2).encode("utf-8"))
    return manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("result", nargs="?", help="Comparison JSON")
    parser.add_argument("--source-docx")
    parser.add_argument("--output")
    parser.add_argument("--mapping", help="Explicit paragraph/table mapping JSON")
    parser.add_argument("--reconstructed", action="store_true")
    parser.add_argument("--inspect", metavar="DOCX", help="Print read-only paragraph mapping inventory")
    args = parser.parse_args()
    try:
        if args.inspect:
            value = inspect_docx(args.inspect)
        else:
            if not args.result or not args.output:
                parser.error("result and --output are required unless --inspect is used")
            result = json.loads(Path(args.result).read_text(encoding="utf-8-sig"))
            mapping = json.loads(Path(args.mapping).read_text(encoding="utf-8-sig")) if args.mapping else None
            value = write_redline(result, args.source_docx, args.output, mapping, args.reconstructed)
        print(json.dumps(value, ensure_ascii=False, indent=2))
    except (ValueError, OSError) as exc:
        parser.exit(2, f"Word export failed: {exc}\n")


if __name__ == "__main__":
    main()
