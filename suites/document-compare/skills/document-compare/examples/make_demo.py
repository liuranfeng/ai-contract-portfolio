"""Build explicitly synthetic fixtures and run both comparison modes. No OCR calls."""
import hashlib
import html
import json
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt
ROOT=next(p for p in Path(__file__).resolve().parents if (p/'document-compare-plugin/skills/document-compare/scripts').is_dir() or (p/'skills/document-compare/scripts').is_dir())
PLUGIN=ROOT/'document-compare-plugin' if (ROOT/'document-compare-plugin').is_dir() else ROOT
sys.path.insert(0,str(PLUGIN/'skills/document-compare/scripts'))
from input_adapter import import_document
from alignment import align_documents
from semantics import apply_decisions
from compare import validate_result,write

def build(out):
    out=Path(out);out.mkdir(parents=True,exist_ok=False)
    paragraphs=[
        '模拟维护服务协议',
        '第一条 服务范围',
        '乙方负责本项目设备的日常维护。',
        '第二条 质量保证',
        '设备免费质保期为3年，自验收合格之日起计算。',
        '第三条 服务费用',
        '服务费用每月支付一次。',
        '第四条 服务响应',
        '乙方接到故障通知后2小时内到达现场。',
        '第五条 材料提交',
        '乙方每季度提交纸质服务报告。',
        '第六条 培训安排',
        '培训由双方指定联系人安排。',
        '第七条 设备清单',
        '第八条 其他约定',
        '未尽事项由双方另行协商，设备标识为😀。'
    ]
    other=paragraphs.copy();other[4]=other[4].replace('3年','5年');other[6]='服务费用按月支付。';other[8]='市区内，乙方接到故障通知后1小时内到达现场。';other.remove(paragraphs[10]);other.insert(other.index(paragraphs[12]),'乙方每年提供两次免费培训。')
    normalized=[]
    for side,lines in enumerate((paragraphs,other)):
        document=Document();normal=document.styles['Normal'];normal.font.name='Microsoft YaHei';normal.font.size=Pt(11)
        normal.element.get_or_add_rPr().rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia','Microsoft YaHei')
        document.sections[0].header.paragraphs[0].text='模拟数据 · 用于套件验证，不是实际合同'
        document.sections[0].footer.paragraphs[0].text='文档比对套件功能演示'
        layouts=[];order=0
        for i,text in enumerate(lines):
            paragraph=document.add_paragraph(text,style='Title' if i==0 else 'Heading 2' if text.startswith('第') else 'Normal')
            if text.startswith('设备免费质保'):paragraph.runs[0].bold=True
            layouts.append({'layout_id':f'L{900-order*7}', 'reading_order':order, 'type':'text','text':text});order+=1
            if text=='第七条 设备清单':
                rows=[['项目','数量','单位'],['控制器','12' if not side else '15','台'],['传感器','30','件']]
                table=document.add_table(rows=3,cols=3);table.style='Table Grid'
                for ri,row in enumerate(rows):
                    for ci,value in enumerate(row):table.cell(ri,ci).text=value
                table_html='<table>'+''.join('<tr>'+''.join('<td>'+html.escape(x)+'</td>' for x in row)+'</tr>' for row in rows)+'</table>'
                layouts.append({'layout_id':f'L{900-order*7}','reading_order':order,'type':'table','text':table_html});order+=1
        name='original.docx' if side==0 else 'revised.docx';path=out/name;document.save(path)
        raw={'schema_version':'fixture','fixture':True,'parse_status':'success','parse_run_id':f'synthetic-fixture-{side}',
             'source_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'pages':[{'page_num':0,'layouts':layouts}]}
        raw_path=out/('left-parsed.json' if side==0 else 'right-parsed.json');write(raw_path,raw)
        d=import_document(raw,side,'模拟原件' if side==0 else '模拟比对件',source=path,raw_bytes=raw_path.read_bytes());normalized.append(d);write(out/('left.json' if side==0 else 'right.json'),d)
    chars=align_documents(*normalized);validate_result(chars);write(out/'character.json',chars)
    pending=align_documents(*normalized,mode='semantic');write(out/'semantic-pending.json',pending)
    decisions=[]
    for row in pending['alignments']:
        if row['operation']=='equal':continue
        a,b=row['leftText'],row['rightText'];status='changed'
        if '服务费用每月支付一次' in a and '服务费用按月支付' in b:status='equivalent';reason='两侧均规定按月付款，付款周期一致，只调整表达。'
        elif not a:reason='比对件增加了每年两次免费培训的明确义务。'
        elif not b:reason='比对件删除了每季度提交纸质服务报告的要求。'
        elif '质保' in a:reason='免费质保由3年延长至5年，起算点仍为验收合格日。'
        elif '到达现场' in a:reason='到场时间由2小时缩短至1小时，同时增加市区范围限定；范围外的要求无法从该句确定。'
        elif '控制器' in a:reason='控制器数量由12台改为15台，传感器数量不变。'
        else:status='uncertain';reason='模拟数据该项需要人工核对对齐和上下文。'
        if row['alignmentStatus']=='unresolved':status='uncertain';reason='当前对齐未确认，不能给出确定语义结论。'
        decisions.append({'id':row['id'],'status':status,'reason':reason,'leftQuote':a,'rightQuote':b})
    payload={'comparisonId':pending['comparisonId'],'fixture':True,'reviewer':'本次构建的模拟用例人工判定','decisions':decisions}
    write(out/'decisions.json',payload);reviewed=apply_decisions(pending,payload);validate_result(reviewed);write(out/'semantic.json',reviewed)
    print(json.dumps({'path':str(out.resolve()),'character':chars['summary'],'semantic':reviewed['summary'],'status':reviewed['status']},ensure_ascii=False))

if __name__=='__main__':build(sys.argv[1])
